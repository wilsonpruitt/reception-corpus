#!/usr/bin/env python3.11
"""Build the full RCL spine (Years A/B/C) from the Vanderbilt Excel downloads.

Source-first: the per-year .xlsx are cached raw in cache/rcl_year_{A,B,C}.xlsx
(Vanderbilt Divinity Library, lectionary.library.vanderbilt.edu — the canonical
machine-readable RCL; citation lists are facts, not copyrightable).

Fidelity contract:
  * refDisplay   = VERBATIM citation text from Vanderbilt (authoritative, auditable;
                   never model-recalled).
  * refKey(s)    = DERIVED OSIS join key(s), KJV/WEB versification, via osis.build_refkey.
                   Every refKey is validated downstream (validate_rcl.py); unparseable
                   citations are surfaced, never silently invented.

Column semantics (uniform across years; Year A's header labels are legacy/mislabeled
but the data columns align):
  0 liturgical date | 1 calendar date | 2 first | 3 psalm | 4 second | 5 gospel

Dual tracks (Season after Pentecost only): col 2 = semicontinuous "FIRST and PSALM",
col 3 = complementary "FIRST and PSALM". Elsewhere col 2 = first, col 3 = psalm.
"""
from __future__ import annotations
import json, re, sys
from pathlib import Path
import openpyxl
import osis

ROOT = Path(__file__).resolve().parent.parent
CACHE = ROOT / "cache"
OUT = ROOT / "data" / "rcl.json"

YEAR_FILES = {"A": "rcl_year_A.xlsx", "B": "rcl_year_B.xlsx", "C": "rcl_year_C.xlsx"}

# ---- book name -> OSIS code -------------------------------------------------
NAME2CODE = {v: k for k, v in osis.OSIS_NAME.items()}
NAME2CODE["Psalm"] = "Ps"          # Vanderbilt writes the singular
NAME2CODE["Psalms"] = "Ps"
# source spelling variants / typos in the Vanderbilt files
NAME2CODE["2 Corinthian"] = "2Cor"
NAME2CODE["1 Corinthian"] = "1Cor"
NAME2CODE["Song of Songs"] = "Song"
# match longest names first so "1 John" beats "John", "Song of Solomon" is whole, etc.
NAMES_SORTED = sorted(NAME2CODE, key=len, reverse=True)


def match_book(cite: str):
    """Strip a leading book name; return (osis_code, remaining_locator)."""
    cite = cite.strip()
    for nm in NAMES_SORTED:
        if cite == nm or cite.startswith(nm + " ") or cite.startswith(nm + ":"):
            return NAME2CODE[nm], cite[len(nm):].strip()
    raise ValueError(f"unrecognized book in citation: {cite!r}")


def _vnum(tok: str) -> int:
    """Verse/chapter number, dropping partial-verse letters (35b -> 35, 9a -> 9)."""
    m = re.match(r"\s*(\d+)", tok)
    if not m:
        raise ValueError(f"no number in token {tok!r}")
    return int(m.group(1))


def parse_locator(loc: str):
    """Locator -> list of spans (c1, v1, c2, v2); v=0 means whole-chapter.

    Handles: whole chapter(s) ('122', '42, 43', '12-14'); verse lists with ',';
    multi-chapter with ';'; cross-chapter ranges ('1:8-2:10'); parentheticals
    (optional verses, folded in); partial-verse letters."""
    loc = loc.strip()
    if not loc:
        raise ValueError("empty locator")

    # Fold parentheses into ordinary, comma-separated segments.
    #   "1:1-9a (9b-12)" -> "1:1-9a, 9b-12"   "2:(1-7), 8" -> "2:, 1-7, 8"
    loc = re.sub(r"\s*\(", ", ", loc).replace(")", "")

    # No ':' anywhere -> pure chapter reference(s).
    if ":" not in loc:
        spans = []
        for seg in re.split(r"[;,]", loc):
            seg = seg.strip()
            if not seg:
                continue
            if "-" in seg:
                a, b = seg.split("-", 1)
                spans.append((_vnum(a), 0, _vnum(b), 0))
            else:
                c = _vnum(seg)
                spans.append((c, 0, c, 0))
        return spans

    # Verse mode: walk ';'/',' segments, tracking the current chapter.
    spans = []
    cur = None
    for seg in re.split(r"[;,]", loc):
        seg = seg.strip()
        if not seg:
            continue
        # A leading "chapter:" sets the current chapter — but only when the colon
        # precedes any "-" (else the colon belongs to a cross-chapter range END,
        # e.g. "22-22:5" = v22 of cur chapter through chapter 22 verse 5).
        if ":" in seg and ("-" not in seg or seg.index(":") < seg.index("-")):
            cpart, _, seg = seg.partition(":")
            cur = int(cpart.strip())
            seg = seg.strip()
            if not seg:                      # e.g. "2:" produced by paren-folding
                continue
        if cur is None:
            raise ValueError(f"verse segment before any chapter in {loc!r}")
        if "-" in seg:
            a, b = seg.split("-", 1)
            v1 = _vnum(a)
            if ":" in b:                     # cross-chapter end "10:8"
                ec, ev = b.split(":")
                ec = int(ec.strip())
                spans.append((cur, v1, ec, _vnum(ev)))
                cur = ec
            else:
                spans.append((cur, v1, cur, _vnum(b)))
        else:
            v = _vnum(seg)
            spans.append((cur, v, cur, v))
    return spans


def span_to_refkey(code, c1, v1, c2, v2) -> str:
    if v1 == 0 and v2 == 0:
        return osis.build_refkey(code, c1) if c1 == c2 else osis.build_refkey(code, c1, None, c2, None)
    if c1 == c2 and v1 == v2:
        return osis.build_refkey(code, c1, v1)
    return osis.build_refkey(code, c1, v1, c2, v2)


# known transcription typos in the Vanderbilt source — fixed in refDisplay only
# (book name is unambiguous from the matched OSIS code; locator stays verbatim)
DISPLAY_FIX = {
    r"^2 Corinthian (?=\d)": "2 Corinthians ",
    r"^1 Corinthian (?=\d)": "1 Corinthians ",
}


def parse_citation(text: str) -> dict:
    """One citation (book + locator, no 'or'/'and') -> reading dict fields."""
    text = text.strip()
    code, loc = match_book(text)
    spans = parse_locator(loc)
    refkeys = [span_to_refkey(code, *s) for s in spans]
    disp = text
    for pat, repl in DISPLAY_FIX.items():
        disp = re.sub(pat, repl, disp)
    out = {"refKey": refkeys[0], "refDisplay": disp}
    if len(refkeys) > 1:
        out["refKeys"] = refkeys
        out["multi"] = True
    return out


ERRORS = []  # (context, citation, message) — collected, never silently dropped


def readings_from_cell(text: str, role: str, track=None, note=None, ctx=""):
    """A reading cell for one role, possibly 'A or B or C' -> list of reading dicts
    (first = primary, rest tagged alt:true)."""
    out = []
    for alt in re.split(r"\bor\b", text):
        alt = alt.strip()
        if not alt:
            continue
        r = {"role": role}
        try:
            r.update(parse_citation(alt))
        except Exception as e:                       # collect, keep building
            ERRORS.append((ctx, alt, str(e)))
            r.update({"refKey": None, "refDisplay": alt, "parseError": str(e)})
        if track:
            r["track"] = track
        if note:
            r["note"] = note
        if out:
            r["alt"] = True
        out.append(r)
    return out


def parse_combined(cell: str, track, note=None, ctx=""):
    """'FIRST [or FIRST2] and PSALM [or PSALM2]' -> first + psalm readings.
    Splits on the FIRST ' and ' (the single first-reading / responsorial pivot)."""
    first_part, _, resp_part = cell.partition(" and ")
    out = readings_from_cell(first_part, "first", track, note, ctx)
    out += readings_from_cell(resp_part, "psalm", track, note, ctx)
    return out


SEASON_RULES = [
    ("Advent", "Advent"),
    ("Christmas", "Christmas"), ("Nativity", "Christmas"), ("Holy Name", "Christmas"),
    ("New Year", "Christmas"),
    ("Epiphany", "Epiphany"), ("Baptism of the Lord", "Epiphany"),
    ("Presentation", "Epiphany"), ("Transfiguration", "Epiphany"),
    ("Ash Wednesday", "Lent"), ("Lent", "Lent"), ("Annunciation", "Lent"),
    ("Holy Week", "Holy Week"), ("Palms", "Holy Week"), ("Passion", "Holy Week"),
    ("Maundy Thursday", "Holy Week"), ("Good Friday", "Holy Week"),
    ("Holy Saturday", "Holy Week"),
    ("Easter Vigil", "Easter"), ("Resurrection", "Easter"), ("Easter", "Easter"),
    ("Ascension", "Easter"),
    ("Day of Pentecost", "Pentecost"),
    ("Visitation", "Easter"), ("Holy Cross", "Season after Pentecost"),
    ("Trinity", "Season after Pentecost"), ("Proper", "Season after Pentecost"),
    ("Reign of Christ", "Season after Pentecost"), ("All Saints", "Season after Pentecost"),
    ("Thanksgiving", "Season after Pentecost"),
]


def season_for(name: str) -> str:
    for needle, season in SEASON_RULES:
        if needle in name:
            return season
    return "Other"


def slugify(name: str, year: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    return f"{s}-{year.lower()}"


def fix_palms(readings):
    """Vanderbilt files the Liturgy-of-the-Palms processional Gospel in the second
    column. If there is no gospel but a 'second' reading is a Gospel book, relabel it."""
    gospels = {"Matt", "Mark", "Luke", "John"}
    has_gospel = any(r["role"] == "gospel" for r in readings)
    if has_gospel:
        return
    for r in readings:
        if r["role"] == "second" and r["refKey"].split(".")[0] in gospels:
            r["role"] = "gospel"


def build_year(year: str):
    wb = openpyxl.load_workbook(CACHE / YEAR_FILES[year], read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    occasions = []
    seen_ids = {}
    for row in ws.iter_rows(values_only=True):
        c = [("" if x is None else str(x)).replace("\n", " ").strip() for x in (list(row) + [""] * 6)[:6]]
        name = c[0]
        if (not name or name.startswith("Revised") or name.startswith("Scripture")
                or "vanderbilt" in name.lower() or name == "Liturgical Date"):
            continue
        colC, colD, colE, colF = c[2], c[3], c[4], c[5]
        ctx = f"{year}/{name}"
        season = season_for(name)
        readings = []

        if " / " in colC:                                    # Easter Vigil sequence
            for idx, seg in enumerate(colC.split(" / "), 1):
                readings += parse_combined(seg.strip(), None, note=f"Easter Vigil reading {idx}", ctx=ctx)
        elif " and " in colC:                                # combined first+psalm cell
            if colD and " and " in colD:                     # dual track
                readings += parse_combined(colC, "semicontinuous", ctx=ctx)
                readings += parse_combined(colD, "complementary", ctx=ctx)
            else:
                readings += parse_combined(colC, None, ctx=ctx)
                if colD:                                     # safety: unexpected D content
                    readings += readings_from_cell(colD, "psalm", ctx=ctx)
        elif (season == "Season after Pentecost"
              and len(re.split(r"\bor\b", colC)) == 2 and len(re.split(r"\bor\b", colD)) == 2):
            # Ordinary-time row given in the inline "Track1 or Track2" form rather than
            # dual columns (e.g. Proper 4 B). Per RCL convention the first option is the
            # semicontinuous track, the second the complementary track (not an alternative).
            cf = [x.strip() for x in re.split(r"\bor\b", colC)]
            cp = [x.strip() for x in re.split(r"\bor\b", colD)]
            for trk, f_, p_ in (("semicontinuous", cf[0], cp[0]), ("complementary", cf[1], cp[1])):
                readings += readings_from_cell(f_, "first", trk, ctx=ctx)
                readings += readings_from_cell(p_, "psalm", trk, ctx=ctx)
        else:                                                # separate first / psalm
            if colC:
                readings += readings_from_cell(colC, "first", ctx=ctx)
            if colD:
                readings += readings_from_cell(colD, "psalm", ctx=ctx)

        if colE:
            if " and " in colE:                              # vigil epistle "Romans … and Psalm 114"
                sec, _, resp = colE.partition(" and ")
                readings += readings_from_cell(sec, "second", ctx=ctx)
                readings += readings_from_cell(resp, "psalm", ctx=ctx)
            else:
                readings += readings_from_cell(colE, "second", ctx=ctx)
        if colF:
            readings += readings_from_cell(colF, "gospel", ctx=ctx)

        fix_palms(readings)

        oid = slugify(name, year)
        if oid in seen_ids:
            seen_ids[oid] += 1
            oid = f"{oid}-{seen_ids[oid]}"
        else:
            seen_ids[oid] = 1
        occasions.append({
            "id": oid, "year": year, "season": season_for(name), "name": name,
            "readings": readings,
        })
    wb.close()
    return occasions


def main():
    all_occ = []
    for yr in ["A", "B", "C"]:
        occ = build_year(yr)
        print(f"Year {yr}: {len(occ)} occasions, "
              f"{sum(len(o['readings']) for o in occ)} readings")
        all_occ += occ
    spine = {
        "dataset": "Revised Common Lectionary — spine (Years A, B, C)",
        "description": (
            "Canonical lectionary calendar in the Wroot data-repository standard. Every reading "
            "carries an OSIS refKey (KJV/WEB versification) — the universal join into the reception "
            "store, Catena, and Lectern's series logic. refDisplay is verbatim from Vanderbilt "
            "(lectionary.library.vanderbilt.edu, per-year Excel downloads cached in cache/); refKeys "
            "are derived and validated. RCL cites NRSV versification — refKeys normalized to KJV/WEB; "
            "Apocrypha refKeys (Wis/Sir/Bar) carry no reception text but keep authoritative refDisplay."
        ),
        "source": "Vanderbilt Divinity Library — Revised Common Lectionary (Excel per year)",
        "refKeyScheme": "OSIS osisRef",
        "versification": "KJV/WEB",
        "tracks": ("Advent through Pentecost use a single reading set (no track field). The Season "
                   "after Pentecost carries dual OT tracks: track='semicontinuous' (Vanderbilt col 3) "
                   "and track='complementary' (col 4); the second reading and Gospel are shared."),
        "readingFields": ("role (first|psalm|second|gospel); refKey + refDisplay; refKeys[]+multi for "
                          "split verse-spans; track for Pentecost-season OT; alt:true for an 'or' "
                          "alternative to the preceding same-role reading; note for Vigil index."),
        "occasions": all_occ,
    }
    OUT.write_text(json.dumps(spine, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"\nWrote {OUT} — {len(all_occ)} occasions total")
    if ERRORS:
        print(f"\n!! {len(ERRORS)} citation parse error(s):")
        for ctx, cite, msg in ERRORS:
            print(f"   [{ctx}] {cite!r}: {msg}")
        sys.exit(1)
    print("All citations parsed.")


if __name__ == "__main__":
    main()
